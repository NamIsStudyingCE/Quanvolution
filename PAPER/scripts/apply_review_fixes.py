# -*- coding: utf-8 -*-
"""apply_review_fixes.py — áp dụng các fix sau review slide ChinhThuc:
A) Deck pptx: bảng 9 (ddof=1), bảng 10 (thay số ma bằng canonical + thêm hàng
   Quán quân), wording slide 1/4, ref [6] slide 13.
B) ĐỒNG BỘ THS. + tên/MSSV vào KLTN_draft_full / skeleton / de_cuong_v2."""
import json
from pathlib import Path
from copy import deepcopy
from pptx import Presentation
from pptx.util import Inches
from docx import Document

ROOT = Path('.').resolve()
CANON = json.load(open(ROOT / 'results' / 'reconciliation_canonical.json', encoding='utf-8'))
B, O = CANON['breastmnist']['models'], CANON['octmnist']['models']
MET = ['acc', 'bacc', 'f1', 'mcc', 'auc', 'pr_auc']

# mean -> canonical std ddof=1 (đã kiểm chứng 0 collision)
mean2std = {}
for M in (B, O):
    for m, mm in M.items():
        for k in MET:
            mean2std.setdefault(round(mm[k]['mean'], 4), mm[k]['std'])

DECK = ROOT / 'slides' / 'KLTN_BaoCao_KhoaKTMT_ChinhThuc.pptx'
prs = Presentation(str(DECK))

def iter_texts(shape):
    if shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                yield r
    if shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                for p in cell.text_frame.paragraphs:
                    for r in p.runs:
                        yield r

def set_cell(cell, text):
    """Ghi text vào cell, giữ format của run đầu tiên."""
    tf = cell.text_frame
    paras = tf.paragraphs
    if paras and paras[0].runs:
        paras[0].runs[0].text = text
        for r in paras[0].runs[1:]:
            r.text = ''
        for p in paras[1:]:
            for r in p.runs:
                r.text = ''
    else:
        cell.text = text

# ---------- A1. Slide 9: chuẩn hóa std theo mean-lookup ----------
import re
s9 = prs.slides[8]
n_fix = 0
for sh in s9.shapes:
    if not sh.has_table:
        continue
    for row in sh.table.rows:
        for cell in row.cells:
            m = re.search(r'(\d\.\d{4})\s*±\s*(\d\.\d{4})', cell.text)
            if not m:
                continue
            mean = round(float(m.group(1)), 4)
            if mean in mean2std and abs(float(m.group(2)) - mean2std[mean]) > 5e-5:
                new_text = cell.text.replace(f"± {m.group(2)}", f"± {mean2std[mean]:.4f}")
                set_cell(cell, new_text)
                n_fix += 1
print(f'A1 slide 9: {n_fix} ô std chuẩn hóa ddof=1')

# ---------- A2. Slide 10: thay bảng bằng canonical + thêm hàng Quán quân ----------
s10 = prs.slides[9]
CANON_OCT = {
    'classical_cnn':        dict(auc='0.7505 ± 0.0240', pr='0.4991 ± 0.0297', acc='44.33% (Cao nhất)', kernel='20 tham số'),
    'trainable_strongly':   dict(auc='0.6922 ± 0.0199', pr='0.4365 ± 0.0289', acc='40.20%', kernel='24 tham số'),
    'fixed_strongly':       dict(auc='0.6690 ± 0.0055', pr='0.4175 ± 0.0047', acc='40.34%', kernel='0 tham số (khóa)'),
    'trainable_basic':      dict(auc='0.6704 ± 0.0106', pr='0.4102 ± 0.0131', acc='39.55%', kernel='8 tham số'),
    'fixed_basic':          dict(auc='0.6711 ± 0.0042', pr='0.4186 ± 0.0074', acc='40.75%', kernel='0 tham số (khóa)'),
    'fixed_champion_gd2':   dict(auc='0.6912 ± 0.0071', pr='0.4443 ± 0.0088', acc='40.48%', kernel='0 tham số (khóa)'),
}
ROW_LABEL = {
    'classical_cnn': 'CNN Cổ Điển (Chuẩn)',
    'trainable_strongly': 'Mạch Lượng Tử Strongly (Tự học)',
    'fixed_strongly': 'Mạch Lượng Tử Strongly (Tĩnh)',
    'trainable_basic': 'Mạch Lượng Tử Basic (Tự học)',
    'fixed_basic': 'Mạch Lượng Tử Basic (Tĩnh)',
    'fixed_champion_gd2': 'Quán quân Tĩnh GĐ2 (random_L1)',
}
fixed_rows = 0
for sh in s10.shapes:
    if not sh.has_table:
        continue
    tbl = sh.table
    for row in tbl.rows:
        label = row.cells[0].text
        for key, disp in ROW_LABEL.items():
            if disp in label or label.strip() == disp:
                v = CANON_OCT[key]
                set_cell(row.cells[2], v['auc'])
                set_cell(row.cells[3], v['pr'])
                set_cell(row.cells[4], v['acc'])
                fixed_rows += 1
    # thêm hàng Quán quân nếu thiếu
    if not any('random_L1' in r.cells[0].text for r in tbl.rows):
        trs = tbl._tbl.findall('{http://schemas.openxmlformats.org/drawingml/2006/main}tr')
        new_tr = deepcopy(trs[3])  # copy hàng Fixed Strongly làm khuôn
        tbl._tbl.append(new_tr)
        new_row = tbl.rows[len(tbl.rows) - 1]
        v = CANON_OCT['fixed_champion_gd2']
        vals = ['Quán quân GĐ2 (random_L1)', '0 (khóa)', v['auc'], v['pr'], v['acc']]
        set_cell_cells = list(new_row.cells)
        for j, val in enumerate(vals):
            set_cell_cells_j = set_cell_cells[j]
            tf = set_cell_cells_j.text_frame
            if tf.paragraphs and tf.paragraphs[0].runs:
                tf.paragraphs[0].runs[0].text = val
                for r in tf.paragraphs[0].runs[1:]:
                    r.text = ''
            else:
                set_cell_cells_j.text = val
        fixed_rows += 1
print(f'A2 slide 10: {fixed_rows} hàng chuẩn hóa/thêm')

# ---------- A3. Wording slide 1 & 4 ----------
n_w = 0
for sh in prs.slides[0].shapes:
    if sh.has_text_frame and 'BÌNH DUYỆT' in sh.text_frame.text and 'ĐANG' not in sh.text_frame.text:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if 'BÌNH DUYỆT' in r.text:
                    r.text = r.text.replace('ĐÃ NỘP & BÌNH DUYỆT', 'ĐÃ NỘP & ĐANG BÌNH DUYỆT')
                    n_w += 1
for sh in prs.slides[3].shapes:
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if 'loại trừ hoàn toàn yếu tố ngẫu nhiên' in r.text:
                    r.text = r.text.replace('loại trừ hoàn toàn yếu tố ngẫu nhiên',
                                            'loại trừ khả năng kết luận chỉ dựa trên may rủi')
                    n_w += 1
print(f'A3 wording fixes: {n_w}')

# ---------- A4. Slide 13: sửa tên bài báo ref Matondo ----------
n_r = 0
for sh in prs.slides[12].shapes:
    if not sh.has_text_frame:
        continue
    for p in sh.text_frame.paragraphs:
        for r in p.runs:
            if 'Hybrid Quantum-Classical Convolutional Neural Networks' in r.text:
                r.text = r.text.replace('Hybrid Quantum-Classical Convolutional Neural Networks',
                                        'Quanvolutional Neural Networks')
                n_r += 1
            elif 'Hybrid Quantum-Classical' in r.text and 'Breast Cancer' in r.text:
                r.text = r.text.replace('Hybrid Quantum-Classical Convolutional Neural Networks',
                                        'Quanvolutional Neural Networks')
                n_r += 1
print(f'A4 ref [6] title fixes: {n_r}')

prs.save(str(DECK))
print('deck saved:', DECK.name)

# ============ B. ĐỒNG BỘ DOCX LUẬN VĂN ============
print('--- B. docx sync ---')
def fix_docx(path):
    p = Path(path)
    if not p.exists():
        print(f'  {p.name}: MISSING'); return
    d = Document(str(p))
    n = 0
    for para in d.paragraphs:
        for r in para.runs:
            if 'TS. Nguyen Duy Xuan Bach' in r.text:
                r.text = r.text.replace('TS. Nguyen Duy Xuan Bach', 'ThS. Nguyễn Duy Xuân Bách')
                n += 1
            if 'TS. Nguyễn Duy Xuân Bách' in r.text:
                r.text = r.text.replace('TS. Nguyễn Duy Xuân Bách', 'ThS. Nguyễn Duy Xuân Bách')
                n += 1
            if '<HỌ TÊN SINH VIÊN> – <MSSV>' in r.text:
                r.text = r.text.replace('<HỌ TÊN SINH VIÊN> – <MSSV>', 'NGUYỄN HOÀNG NAM – 22520916')
                n += 1
    # bảng: bìa phụ có thể nằm trong table cell
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for r in para.runs:
                        if 'TS. Nguyen Duy Xuan Bach' in r.text or 'TS. Nguyễn Duy Xuân Bách' in r.text:
                            r.text = r.text.replace('TS. Nguyen Duy Xuan Bach', 'ThS. Nguyễn Duy Xuân Bách') \
                                           .replace('TS. Nguyễn Duy Xuân Bách', 'ThS. Nguyễn Duy Xuân Bách')
                            n += 1
    try:
        d.save(str(p))
        print(f'  {p.name}: {n} fixes saved')
    except PermissionError:
        print(f'  {p.name}: FILE ĐANG MỞ TRONG WORD — chưa lưu được, hãy đóng file rồi chạy lại')

import re as _re
def fix_decuong(path):
    p = Path(path)
    if not p.exists():
        print(f'  {p.name}: MISSING'); return
    d = Document(str(p))
    n = 0
    def walk(paras):
        nn = 0
        for para in paras:
            for r in para.runs:
                old = r.text
                new = old.replace('TS. Nguyen Duy Xuan Bach', 'ThS. Nguyễn Duy Xuân Bách') \
                         .replace('Cán bộ hướng dẫn:  TS.', 'Cán bộ hướng dẫn: ThS.') \
                         .replace('<Họ tên sinh viên – MSSV>', 'NGUYỄN HOÀNG NAM – 22520916') \
                         .replace('<Họ tên sinh viên 1 – MSSV 1>', 'NGUYỄN HOÀNG NAM – 22520916')
                if new != old:
                    r.text = new; nn += 1
        return nn
    n += walk(d.paragraphs)
    for tbl in d.tables:
        for row in tbl.rows:
            for cell in row.cells:
                n += walk(cell.paragraphs)
    d.save(str(p))
    print(f'  {p.name}: {n} fixes')

fix_docx(ROOT / 'GD4' / 'KLTN_draft_full.docx')
fix_docx(ROOT / 'GD4' / 'KLTN_skeleton.docx')
fix_decuong(ROOT / 'GD4' / 'de_cuong_chi_tiet_filled_v2.docx')
fix_decuong(ROOT / 'GD4' / 'de_cuong_chi_tiet_filled.docx')
