# -*- coding: utf-8 -*-
"""add_source_notes.py — chèn 2 bổ sung đã chốt (KHÔNG đụng gì khác):
1. S9/S10: dòng chú thích nguồn số liệu (10 seeds · ddof=1 · Bảng 4.1/4.2)
2. S6: dòng tham số 3 loại mạch (Basic 4L · Strongly 12L · Random 0)
Vị trí tính để KHÔNG đè lên callout/copyright đã có."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

F = 'slides/KLTN_BaoCao_KhoaKTMT_ChinhThuc.pptx'
prs = Presentation(F)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)

def note(slide, x, y, w, text, size=10.5):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.28))
    tf = tb.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]; r = p.add_run(); r.text = text
    r.font.name = 'Arial'; r.font.size = Pt(size); r.font.color.rgb = MUTED
    return tb

# S9 (index 8): ghi chú nguồn dưới hàng callout, cột trái (tránh copyright ở giữa-dưới)
note(prs.slides[8], 0.78, 6.72, 3.9,
     'Số liệu: 10 seeds · mean ± std (ddof=1) · Nguồn: Bảng 4.1 luận văn')

# S10 (index 9): tương tự
note(prs.slides[9], 0.78, 6.82, 3.9,
     'Số liệu: 10 seeds · mean ± std (ddof=1) · Nguồn: Bảng 4.2 luận văn')

# S6 (index 5): dòng tham số mạch cuối thẻ "Ba Cấu Trúc Mạch"
for sh in prs.slides[5].shapes:
    if sh.has_text_frame and 'Ba Cấu Trúc Mạch Lượng Tử' in sh.text_frame.text:
        tf = sh.text_frame
        p = tf.add_paragraph(); p.space_before = Pt(8)
        r = p.add_run()
        r.text = 'Tham số/tầng: Basic 4L · Strongly 12L · Random 0 (đóng băng)'
        r.font.name = 'Arial'; r.font.size = Pt(13); r.font.bold = True
        r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
        print('S6: đã thêm dòng tham số mạch')
        break

prs.save(F)
print('saved', F)
